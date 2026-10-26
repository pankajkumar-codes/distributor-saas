"""Reads CSV and XLSX bytes into a plain grid of strings.

Every cell comes out as a string (or "" when empty) so the raw source value can be
stored exactly as the marketplace sent it. Typing happens later, in normalize.py,
under rules the mapping profile states explicitly.
"""

import csv
import datetime as dt
import hashlib
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLUMNS = 300
# An .xlsx is a zip; refuse archives that inflate far beyond their upload size.
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


class FileRejected(ValueError):
    """The whole file is unusable; no rows were read."""


@dataclass(frozen=True)
class SourceTable:
    filename: str
    checksum_sha256: str
    sheet: str | None
    # rows[i] is source line/row i + 1, exactly as numbered in a spreadsheet app.
    rows: list[list[str]]


def checksum(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def read_table(content: bytes, filename: str, *, sheet: str | None = None) -> SourceTable:
    if len(content) == 0:
        raise FileRejected("file is empty")
    if len(content) > MAX_FILE_BYTES:
        raise FileRejected(f"file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")

    name = filename.lower()
    if name.endswith((".csv", ".txt")):
        if sheet is not None:
            raise FileRejected("a sheet name only applies to Excel files")
        rows, used_sheet = _read_csv(content), None
    elif name.endswith(".xlsx"):
        rows, used_sheet = _read_xlsx(content, sheet)
    elif name.endswith(".xls"):
        raise FileRejected("legacy .xls is not supported; save the report as .xlsx or .csv")
    else:
        raise FileRejected("unsupported file type; upload .csv or .xlsx")

    return SourceTable(filename=filename, checksum_sha256=checksum(content), sheet=used_sheet, rows=rows)


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        # utf-8-sig drops the byte-order mark Excel adds when saving "CSV UTF-8".
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        # Guessing a legacy code page could silently corrupt names and addresses.
        raise FileRejected(f"file is not UTF-8 text (byte {exc.start}); re-save it as CSV UTF-8") from exc
    if "\x00" in text:
        raise FileRejected("file contains NUL bytes; it is probably not a CSV")

    rows: list[list[str]] = []
    try:
        for row in csv.reader(io.StringIO(text, newline="")):
            _check_shape(len(rows) + 1, len(row))
            rows.append(row)
    except csv.Error as exc:
        raise FileRejected(f"malformed CSV near line {len(rows) + 1}: {exc}") from exc
    return rows


def _read_xlsx(content: bytes, sheet: str | None) -> tuple[list[list[str]], str]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            inflated = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise FileRejected("file is not a valid .xlsx workbook") from exc
    if inflated > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise FileRejected("workbook expands to more than the allowed size")

    import openpyxl  # imported lazily so CSV-only paths do not pay for it

    try:
        # data_only=False: a formula cell reads as its "=..." text, so normalize.py
        # rejects it instead of trusting a cached result or a silent blank.
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises a variety of types for corrupt files
        raise FileRejected("file is not a readable .xlsx workbook") from exc

    try:
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                # Picking one silently could import the wrong report.
                raise FileRejected(
                    "workbook has several sheets; choose one of: " + ", ".join(workbook.sheetnames)
                )
            sheet = workbook.sheetnames[0]
        elif sheet not in workbook.sheetnames:
            raise FileRejected(f"sheet {sheet!r} not found; available: " + ", ".join(workbook.sheetnames))

        worksheet = workbook[sheet]
        # Read-only mode trusts the sheet's stored <dimension>, and some marketplace
        # workbooks (Meesho, Flipkart) have none or a wrong one, so every row after the
        # first would be lost. Re-measure from the cells instead (LLD-7 L7-1, RM47-J16).
        worksheet.reset_dimensions()
        rows: list[list[str]] = []
        for values in worksheet.iter_rows(values_only=True):
            _check_shape(len(rows) + 1, len(values))
            rows.append([_cell_text(v) for v in values])
    finally:
        workbook.close()

    # Trailing empty rows are noise.
    while rows and not any(rows[-1]):
        rows.pop()
    return rows, sheet


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # Excel stores numbers as binary doubles. str() gives the shortest decimal that
        # round-trips, which is what Excel displays for 199.99; Decimal(value) would not.
        text = format(Decimal(repr(value)), "f")
        return text.removesuffix(".0")
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _check_shape(row_number: int, width: int) -> None:
    if row_number > MAX_ROWS:
        raise FileRejected(f"file has more than {MAX_ROWS} rows; split it and upload each part")
    if width > MAX_COLUMNS:
        raise FileRejected(f"row {row_number} has more than {MAX_COLUMNS} columns")
